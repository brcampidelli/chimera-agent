"""Render a parsed :class:`DocumentSpec` to .docx, .xlsx or .pptx bytes — the ``documents-out`` extra.

Each renderer imports its library inside the function, so a machine without the extra can still
import this module, write a PDF, and get an install hint for the rest. Each returns bytes rather than
writing a file: the tool writes once, after the whole document rendered, so a failure half-way leaves
nothing behind instead of a truncated file with the right name.

The one rule here that is a security rule rather than a layout one is in :func:`_set_cell`: text from
the model never becomes a spreadsheet formula.
"""

from __future__ import annotations

import io
import re
from typing import Any

from chimera.tools.document_spec import (
    Bullets,
    Cell,
    DocumentSpec,
    Heading,
    PageBreak,
    Paragraph,
    Table,
)

INSTALL_HINT = (
    "error: writing {fmt} needs the 'documents-out' extra — install with: "
    "pip install 'chimera-agent[documents-out]' (PDF needs no extra)"
)

# --- docx -----------------------------------------------------------------------------------------


def _docx_table(doc: Any, table: Table) -> None:
    ncols = max([len(table.columns)] + [len(r) for r in table.rows]) or 1
    rows = ([table.columns] if table.columns else []) + list(table.rows)
    grid = doc.add_table(rows=len(rows), cols=ncols)
    grid.style = "Table Grid"
    # Filled row by row through `row.cells`, never `grid.cell(r, c)`: that one rebuilds the list of
    # EVERY cell in the table on each call, so a table took O(cells^2) — 400x10 measured at 154 s,
    # and the spec's 5000x50 ceiling would have held the agent's turn for hours.
    for r, (row_obj, row) in enumerate(zip(grid.rows, rows, strict=True)):
        cells = row_obj.cells
        for c in range(ncols):
            value = row[c] if c < len(row) else None
            cells[c].text = "" if value is None else str(value)
            if r == 0 and table.columns:
                for run in cells[c].paragraphs[0].runs:
                    run.bold = True


def render_docx(spec: DocumentSpec) -> bytes:
    from docx import Document  # lazy: the documents-out extra

    doc = Document()
    if spec.title:
        doc.add_heading(spec.title, level=0)
        doc.core_properties.title = spec.title
    for block in spec.blocks:
        if isinstance(block, Heading):
            doc.add_heading(block.text, level=block.level)
        elif isinstance(block, Paragraph):
            doc.add_paragraph(block.text)
        elif isinstance(block, Bullets):
            style = "List Number" if block.numbered else "List Bullet"
            for item in block.items:
                doc.add_paragraph(item, style=style)
        elif isinstance(block, Table):
            if block.name:
                doc.add_paragraph(block.name).runs[0].bold = True
            _docx_table(doc, block)
        elif isinstance(block, PageBreak):
            doc.add_page_break()  # type: ignore[no-untyped-call]  # python-docx leaves this one unannotated
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --- xlsx -----------------------------------------------------------------------------------------

#: What Excel refuses in a sheet name, and its length limit.
_SHEET_FORBIDDEN = re.compile(r"[\[\]:*?/\\]")
_SHEET_MAX = 31


def _sheet_name(raw: str, index: int, taken: set[str]) -> str:
    base = _SHEET_FORBIDDEN.sub("_", raw).strip().strip("'")[:_SHEET_MAX] or f"Sheet{index + 1}"
    name, n = base, 2
    while name.lower() in taken:
        suffix = f" ({n})"
        name = base[: _SHEET_MAX - len(suffix)] + suffix
        n += 1
    taken.add(name.lower())
    return name


def _set_cell(cell: Any, value: Cell) -> None:
    """Write one value. A string is stored as a string — never as a formula.

    openpyxl reads any string that starts with ``=`` as a formula, and Excel and LibreOffice will
    evaluate one on open: ``=HYPERLINK(...)``, ``=WEBSERVICE(...)``, DDE payloads. Model text in a
    cell is untrusted text — it may have come from a fetched page — so the type is forced to text
    and the cell carries Excel's quote prefix, which also keeps a later manual edit of the cell from
    turning it into one. ``+``, ``-`` and ``@`` lead a formula in some spreadsheet programs too, and
    the same two settings cover them. A number from the spec stays a number.
    """
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"
        cell.quotePrefix = True


def xlsx_sheets(spec: DocumentSpec) -> tuple[Table, ...]:
    """The spec's sheets, or — for a spec written as a document — its tables, one sheet each."""
    if spec.sheets:
        return spec.sheets
    return tuple(b for b in spec.blocks if isinstance(b, Table))


def render_xlsx(spec: DocumentSpec) -> bytes:
    from openpyxl import Workbook  # lazy: the documents-out extra
    from openpyxl.styles import Font

    sheets = xlsx_sheets(spec)
    if not sheets:
        raise ValueError("an xlsx needs 'sheets' (or tables among the 'blocks') — there are none")
    wb = Workbook()
    wb.remove(wb.active)
    taken: set[str] = set()
    for index, table in enumerate(sheets):
        ws = wb.create_sheet(_sheet_name(table.name, index, taken))
        row_no = 1
        if table.columns:
            for c, title in enumerate(table.columns, start=1):
                cell = ws.cell(row=1, column=c)
                _set_cell(cell, title)
                cell.font = Font(bold=True)
            row_no = 2
        for row in table.rows:
            for c, value in enumerate(row, start=1):
                _set_cell(ws.cell(row=row_no, column=c), value)
            row_no += 1
    if spec.title:
        wb.properties.title = spec.title
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- pptx -----------------------------------------------------------------------------------------

# Layout indices of the default template python-pptx ships.
_TITLE_SLIDE, _TITLE_AND_CONTENT, _TITLE_ONLY = 0, 1, 5
_EMU_PER_INCH = 914400


def _pptx_table(slide: Any, table: Table) -> None:
    ncols = max([len(table.columns)] + [len(r) for r in table.rows]) or 1
    rows = ([table.columns] if table.columns else []) + list(table.rows)
    left, top = int(0.5 * _EMU_PER_INCH), int(1.5 * _EMU_PER_INCH)
    width, height = int(9 * _EMU_PER_INCH), int(0.4 * _EMU_PER_INCH) * len(rows)
    grid = slide.shapes.add_table(len(rows), ncols, left, top, width, height).table
    for r, row in enumerate(rows):
        for c in range(ncols):
            value = row[c] if c < len(row) else None
            grid.cell(r, c).text = "" if value is None else str(value)


def render_pptx(spec: DocumentSpec) -> bytes:
    from pptx import Presentation  # lazy: the documents-out extra

    if not spec.slides:
        raise ValueError("a pptx needs 'slides' — there are none")
    deck = Presentation()
    if spec.title:
        cover = deck.slides.add_slide(deck.slide_layouts[_TITLE_SLIDE])
        cover.shapes.title.text = spec.title
        deck.core_properties.title = spec.title
    for spec_slide in spec.slides:
        layout = _TITLE_ONLY if spec_slide.table is not None else _TITLE_AND_CONTENT
        slide = deck.slides.add_slide(deck.slide_layouts[layout])
        slide.shapes.title.text = spec_slide.title
        if spec_slide.table is not None:
            _pptx_table(slide, spec_slide.table)
        else:
            body = slide.placeholders[1].text_frame
            for i, bullet in enumerate(spec_slide.bullets):
                paragraph = body.paragraphs[0] if i == 0 else body.add_paragraph()
                paragraph.text = bullet
        if spec_slide.notes:
            slide.notes_slide.notes_text_frame.text = spec_slide.notes
    buf = io.BytesIO()
    deck.save(buf)
    return buf.getvalue()
