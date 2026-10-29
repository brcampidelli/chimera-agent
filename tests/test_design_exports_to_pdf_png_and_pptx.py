"""Chimera Design as PDF, PNG and PPTX (PR D7): what each file holds, and what a machine without the
renderer gets instead of a crash.

Pinned here:

- the PDF has one page per artboard, each at its artboard's size (CSS px at 96 dpi = 0.75 pt);
- a PNG is one artboard at its pixel size, and several artboards are a zip named from ids, not text;
- the Chromium page they are drawn in runs no script and fetches nothing but ``data:``;
- the PPTX puts each element where its frame says, in EMU, scaled uniformly to the slide, and an
  artboard of another aspect is letterboxed rather than stretched;
- without Playwright, without the Chromium binary, or without python-pptx, the exporter refuses with
  the command that installs what is missing (and, for Chromium, its 150 MB), and never raises
  anything else.

The PDF and PNG tests skip, saying why, where Chromium cannot launch (CI ships no browser binary) —
the reason printed is the exporter's own refusal, so a skip and a refusal can never disagree.
"""

from __future__ import annotations

import io
import re
import struct
import sys
import threading
import types
import zipfile
from contextlib import contextmanager
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from typing import Any

import pytest

from chimera.design import export_pdf, export_png, export_pptx
from chimera.design.export_html import render_print_html
from chimera.design.exports import (
    EXPORT_FORMATS,
    PPTX_MEDIA_TYPE,
    ExportFailed,
    ExportUnavailable,
    export_document,
)
from chimera.design.model import DesignDocument, wire
from chimera.design.tokens import default_tokens
from tests.design_fixtures import sample_dict, sample_doc

EMU = 9525


def _render_or_skip(fmt: str, doc: DesignDocument) -> Any:
    try:
        return export_document(doc, fmt)
    except ExportUnavailable as exc:
        pytest.skip(f"chromium unavailable: {exc}")


def _png_size(data: bytes) -> tuple[int, int]:
    assert data[:8] == b"\x89PNG\r\n\x1a\n"
    width, height = struct.unpack(">II", data[16:24])
    return width, height


def _doc(boards: list[dict[str, Any]], nodes: dict[str, Any]) -> DesignDocument:
    return DesignDocument.model_validate({
        "schemaVersion": 2, "id": "d_geo", "title": "Geometry", "revision": 1,
        "updatedAt": "2026-10-18T00:00:00+00:00", "tokens": wire(default_tokens()),
        "pages": [{"id": "p1", "name": "Home", "artboards": boards}],
        "nodes": nodes, "comments": [],
    })


# ------------------------------------------------------------------ PDF


def test_the_pdf_has_one_page_per_artboard_at_the_artboards_size() -> None:
    out = _render_or_skip("pdf", sample_doc())
    assert out.media_type == "application/pdf" and out.extension == "pdf"
    assert out.data.startswith(b"%PDF-")
    boxes = [
        tuple(float(v) for v in m.split())
        for m in re.findall(rb"/MediaBox\s*\[\s*([0-9.\s]+?)\s*\]", out.data)
    ]
    assert len(boxes) == 2, boxes
    # 800x600 px and 390x844 px, at 0.75 pt a px; Chromium rounds to its device grid.
    for (_x0, _y0, w, h), (pw, ph) in zip(boxes, [(800, 600), (390, 844)], strict=True):
        assert abs(w - pw * 0.75) < 1.5 and abs(h - ph * 0.75) < 1.5


def test_a_design_without_artboards_is_refused_before_chromium() -> None:
    doc = _doc([], {})
    for render in (export_pdf.render_pdf, export_pptx.render_pptx):
        with pytest.raises(ExportFailed, match="no artboard"):
            render(doc)
    with pytest.raises(ExportFailed, match="no artboard"):
        export_png.render_png(doc)


# ------------------------------------------------------------------ PNG


def test_one_artboard_is_one_png_at_its_pixel_size() -> None:
    raw = sample_dict()
    raw["pages"] = raw["pages"][:1]
    raw["nodes"].pop("t3")
    out = _render_or_skip("png", DesignDocument.model_validate(raw))
    assert out.media_type == "image/png" and out.extension == "png"
    assert _png_size(out.data) == (800, 600)


def test_several_artboards_are_a_zip_named_from_ids_never_from_text() -> None:
    out = _render_or_skip("png", sample_doc())
    assert out.media_type == "application/zip" and out.extension == "zip"
    with zipfile.ZipFile(io.BytesIO(out.data)) as archive:
        assert archive.namelist() == ["01-a1.png", "02-a2.png"]
        assert _png_size(archive.read("01-a1.png")) == (800, 600)
        assert _png_size(archive.read("02-a2.png")) == (390, 844)


def test_a_png_larger_than_chromium_can_draw_is_refused_with_the_pdf_offered() -> None:
    doc = _doc([{"id": "a1", "width": 20_000, "height": 100, "nodeIds": []}], {})
    with pytest.raises(ExportFailed, match="PDF"):
        export_png.render_png(doc)


# ------------------------------------------------------------------ the page Chromium draws


class _Counter(BaseHTTPRequestHandler):
    hits: list[str] = []

    def do_GET(self) -> None:  # noqa: N802 — the stdlib's name
        type(self).hits.append(self.path)
        self.send_response(200)
        self.send_header("Content-Type", "image/png")
        self.end_headers()

    def log_message(self, *_args: Any) -> None:
        return


def test_the_page_runs_no_script_and_fetches_nothing_but_data_urls() -> None:
    """Belt and braces under the exporter's escaping: a page that DID carry a script and a remote
    image would still run nothing and fetch nothing."""
    server = ThreadingHTTPServer(("127.0.0.1", 0), _Counter)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    _Counter.hits = []
    url = f"http://127.0.0.1:{server.server_address[1]}/beacon.png"
    hostile = (
        "<!doctype html><html><head><title>before</title></head><body>"
        f'<img src="{url}"><link rel="stylesheet" href="{url}">'
        "<script>document.title = 'ran'</script>"
        '<img id="ok" src="data:image/gif;base64,R0lGODlhAQABAAAAACw=">'
        "</body></html>"
    )
    try:
        with export_pdf.static_page(hostile, width=200, height=200) as page:
            title = page.title()
    except ExportUnavailable as exc:
        pytest.skip(f"chromium unavailable: {exc}")
    finally:
        server.shutdown()
    assert title == "before", "a script ran"
    assert _Counter.hits == [], "the page reached the network"


def test_the_print_page_is_the_artboards_alone_with_the_exporters_policy() -> None:
    page = render_print_html(sample_doc())
    assert "<h1" not in page and "<h2" not in page and "<script" not in page.lower()
    assert 'data-chimera-artboard="0"' in page and 'data-chimera-artboard="1"' in page
    assert "@page b0{size:800px 600px;margin:0}" in page
    assert "@page b1{size:390px 844px;margin:0}" in page
    assert "Content-Security-Policy" in page


# ------------------------------------------------------------------ refusals


def test_without_playwright_the_refusal_says_how_to_reinstall(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setitem(sys.modules, "playwright.sync_api", None)
    with pytest.raises(ExportUnavailable, match="pip install --upgrade chimera-agent"):
        export_pdf.render_pdf(sample_doc())


class _FakeChromium:
    def __init__(self, error: Exception) -> None:
        self.error = error

    def launch(self, **_kw: Any) -> Any:
        raise self.error


class _FakeManager:
    stopped = 0

    def __init__(self, error: Exception) -> None:
        self.chromium = _FakeChromium(error)

    def stop(self) -> None:
        type(self).stopped += 1


def _fake_playwright(monkeypatch: pytest.MonkeyPatch, error: Exception) -> None:
    module = types.ModuleType("playwright.sync_api")
    module.sync_playwright = lambda: types.SimpleNamespace(  # type: ignore[attr-defined]
        start=lambda: _FakeManager(error)
    )
    monkeypatch.setitem(sys.modules, "playwright.sync_api", module)
    _FakeManager.stopped = 0


@pytest.mark.parametrize("fmt", ["pdf", "png"])
def test_html_artboards_keep_scripts_disabled_and_network_offline(
    monkeypatch: pytest.MonkeyPatch, fmt: str
) -> None:
    captured: dict[str, Any] = {}

    class FakePage:
        def set_default_timeout(self, _timeout: int) -> None:
            return None

        def set_content(self, _html: str, **_kwargs: Any) -> None:
            return None

        def pdf(self, **_kwargs: Any) -> bytes:
            return b"%PDF-fake"

        def locator(self, _selector: str) -> Any:
            return types.SimpleNamespace(screenshot=lambda **_kw: b"png")

    class FakeContext:
        def route(self, *_args: Any) -> None:
            return None

        def new_page(self) -> FakePage:
            return FakePage()

    class FakeBrowser:
        def new_context(self, **kwargs: Any) -> FakeContext:
            captured.update(kwargs)
            return FakeContext()

        def close(self) -> None:
            return None

    class FakeManager:
        chromium = types.SimpleNamespace(launch=lambda **_kw: FakeBrowser())

        def stop(self) -> None:
            return None

    module = types.ModuleType("playwright.sync_api")
    module.sync_playwright = lambda: types.SimpleNamespace(start=lambda: FakeManager())  # type: ignore[attr-defined]
    monkeypatch.setitem(sys.modules, "playwright.sync_api", module)
    doc = _doc([{"id": "html", "width": 300, "height": 160, "nodeIds": [],
                 "html": "<main><p>offline</p></main>"}], {})
    export_document(doc, fmt)
    assert captured["java_script_enabled"] is False
    assert captured["offline"] is True


@pytest.mark.parametrize("fmt", ["pdf", "png"])
def test_without_chromium_the_refusal_names_the_install_and_its_size(
    monkeypatch: pytest.MonkeyPatch, fmt: str
) -> None:
    _fake_playwright(monkeypatch, RuntimeError(
        "BrowserType.launch: Executable doesn't exist at /x/chrome\n"
        "Please run the following command to download new browsers: playwright install"
    ))
    with pytest.raises(ExportUnavailable) as caught:
        export_document(sample_doc(), fmt)
    message = str(caught.value)
    assert "playwright install chromium" in message and "150 MB" in message
    assert _FakeManager.stopped == 1, "the driver was left running"


def test_a_chromium_that_will_not_start_is_a_refusal_not_a_crash(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _fake_playwright(monkeypatch, OSError("libnss3.so: cannot open shared object file"))
    with pytest.raises(ExportUnavailable, match="--with-deps chromium"):
        export_pdf.render_pdf(sample_doc())


def test_without_python_pptx_the_refusal_names_the_extra(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setitem(sys.modules, "pptx", None)
    with pytest.raises(ExportUnavailable, match="documents-out"):
        export_pptx.render_pptx(sample_doc())


def test_an_unknown_format_is_a_value_error_and_the_list_is_the_four() -> None:
    assert EXPORT_FORMATS == ("html", "pdf", "png", "pptx")
    with pytest.raises(ValueError, match="svg"):
        export_document(sample_doc(), "svg")


# ------------------------------------------------------------------ PPTX

def _slides(data: bytes) -> Any:
    pptx = pytest.importorskip("pptx", reason="python-pptx is the documents-out extra")
    return pptx.Presentation(io.BytesIO(data))


def _by_name(slide: Any) -> dict[str, Any]:
    return {shape.name: shape for shape in slide.shapes}

class _FakeTextBox:
    """One text element as the PPTX exporter measures it in Chromium: its text, box and computed style."""

    def __init__(self, text: str, box: dict[str, float], style: str,
                 error: Exception | None = None) -> None:
        self._text, self._box, self._style, self._error = text, box, style, error

    def inner_text(self) -> str:
        if self._error is not None:
            raise self._error
        return self._text

    def bounding_box(self) -> dict[str, float]:
        return self._box

    def evaluate(self, _script: str) -> str:
        return self._style


def _fake_html_artboard(monkeypatch: pytest.MonkeyPatch, texts: list[Any],
                        *, size: tuple[int, int] = (300, 160)) -> None:
    """Stand in for the locked Chromium page: a white screenshot and the given text elements."""
    from PIL import Image

    png = io.BytesIO()
    Image.new("RGB", (2, 2), "white").save(png, format="PNG")
    picture = png.getvalue()

    class ArtboardLocator:
        def screenshot(self, **_kwargs: Any) -> bytes:
            return picture

        def evaluate(self, _script: str) -> None:
            return None

        def locator(self, _selector: str) -> Any:
            return types.SimpleNamespace(all=lambda: texts)

    class FakePage:
        def locator(self, _selector: str) -> ArtboardLocator:
            return ArtboardLocator()

        def add_style_tag(self, **_kwargs: Any) -> None:
            return None

    @contextmanager
    def fake_page(_html: str, *, width: int, height: int) -> Any:
        assert (width, height) == size
        yield FakePage()

    monkeypatch.setattr(export_pptx, "static_page", fake_page)


def test_html_artboard_pptx_has_background_and_editable_text(monkeypatch: pytest.MonkeyPatch) -> None:
    pptx = pytest.importorskip("pptx", reason="python-pptx is the documents-out extra")
    from PIL import Image

    png = io.BytesIO()
    Image.new("RGB", (2, 2), "white").save(png, format="PNG")

    heading = _FakeTextBox("Editable heading", {"x": 12, "y": 18, "width": 180, "height": 32},
                           "font-size:32px;font-weight:700;color:rgb(0, 0, 0)")
    light = _FakeTextBox("Light text", {"x": 12, "y": 60, "width": 200, "height": 24},
                         'font-size:20px;font-weight:400;color:rgb(255, 255, 255);'
                         'text-align:center;font-family:"Lato", sans-serif')
    _fake_html_artboard(monkeypatch, [heading, light])
    doc = _doc([{"id": "html", "width": 300, "height": 160, "nodeIds": [],
                 "html": "<main><h1>Editable heading</h1></main>"}], {})
    presentation = _slides(export_document(doc, "pptx").data)
    shapes = list(presentation.slides[0].shapes)
    assert any(shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE for shape in shapes)
    editable = [shape for shape in shapes if shape.has_text_frame and shape.text == "Editable heading"]
    assert len(editable) == 1
    shape = editable[0]
    assert (shape.left, shape.top, shape.width, shape.height) == (12 * EMU, 18 * EMU, 180 * EMU, 32 * EMU)
    assert shape.text_frame.paragraphs[0].runs[0].font.size.pt == pytest.approx(24)
    assert str(shape.text_frame.paragraphs[0].runs[0].font.color.rgb) == "000000"

    (white,) = [s for s in shapes if s.has_text_frame and s.text == "Light text"]
    run = white.text_frame.paragraphs[0].runs[0]
    assert str(run.font.color.rgb) == "FFFFFF"  # getComputedStyle's rgb(), not only #RRGGBB
    assert white.text_frame.paragraphs[0].alignment == pptx.enum.text.PP_ALIGN.CENTER
    assert run.font.name == "Lato"  # the first family, without its quotes
    assert run.font.size.pt == pytest.approx(15)  # 20 px at 1:1 is 15 pt


def test_pptx_places_each_element_at_its_frame_in_emu() -> None:
    pytest.importorskip("pptx", reason="python-pptx is the documents-out extra")
    nodes = {
        "r": {"id": "r", "type": "shape", "shape": "rectangle", "frame":
              {"x": 10, "y": 20, "width": 100, "height": 40}, "style": {"fill": "#0F766E"}},
        "o": {"id": "o", "type": "shape", "shape": "ellipse", "frame":
              {"x": 200, "y": 100, "width": 50, "height": 50, "rotation": 30}},
        "t": {"id": "t", "type": "text", "frame": {"x": 40, "y": 300, "width": 300, "height": 60},
              "text": "Launch\nday", "fontToken": "heading", "colorToken": "primary",
              "align": "center"},
        "l": {"id": "l", "type": "shape", "shape": "line", "frame":
              {"x": 400, "y": 400, "width": 200, "height": 100}, "points": [0, 1, 1, 0],
              "style": {"stroke": "#B91C1C", "strokeWidth": 2}},
        "p": {"id": "p", "type": "shape", "shape": "path", "frame":
              {"x": 600, "y": 50, "width": 100, "height": 100}, "points": [0, 0, 1, 0, 0.5, 1]},
        "i": {"id": "i", "type": "image", "frame": {"x": 500, "y": 300, "width": 120, "height": 80},
              "assetId": "missing", "alt": "Team photo"},
    }
    out = export_document(_doc([{"id": "a1", "width": 800, "height": 600, "background": "#F3F5F7",
                                 "nodeIds": list(nodes)}], nodes), "pptx")
    assert out.media_type == PPTX_MEDIA_TYPE and out.extension == "pptx"
    prs = _slides(out.data)
    assert (prs.slide_width, prs.slide_height) == (800 * EMU, 600 * EMU)
    assert len(prs.slides) == 1
    shapes = _by_name(prs.slides[0])

    board = shapes["artboard a1"]
    assert (board.left, board.top, board.width, board.height) == (0, 0, 800 * EMU, 600 * EMU)
    assert str(board.fill.fore_color.rgb) == "F3F5F7"

    rect = shapes["r"]
    assert (rect.left, rect.top, rect.width, rect.height) == (10 * EMU, 20 * EMU, 100 * EMU, 40 * EMU)
    assert str(rect.fill.fore_color.rgb) == "0F766E"
    assert shapes["o"].rotation == pytest.approx(30)

    text = shapes["t"]
    assert (text.left, text.top, text.width, text.height) == (40 * EMU, 300 * EMU, 300 * EMU, 60 * EMU)
    paragraphs = text.text_frame.paragraphs
    assert [p.text for p in paragraphs] == ["Launch", "day"]
    run = paragraphs[0].runs[0]
    assert run.font.size.pt == pytest.approx(24)  # the heading token's 32 px is 24 pt
    assert run.font.bold is True
    assert str(run.font.color.rgb) == "0F766E"
    assert run.font.name == "Inter"

    line = shapes["l"]
    assert (line.begin_x, line.begin_y) == (400 * EMU, 500 * EMU)
    assert (line.end_x, line.end_y) == (600 * EMU, 400 * EMU)
    assert line.line.width == 2 * EMU

    path = shapes["p"]
    assert (path.left, path.top, path.width, path.height) == (600 * EMU, 50 * EMU, 100 * EMU, 100 * EMU)

    image = shapes["i"]
    assert image.text_frame.text == "Team photo", "a missing image is a labelled placeholder"


def test_pptx_letterboxes_an_artboard_of_another_aspect_and_keeps_order() -> None:
    pytest.importorskip("pptx", reason="python-pptx is the documents-out extra")
    nodes = {
        "a": {"id": "a", "type": "shape", "shape": "rectangle",
              "frame": {"x": 0, "y": 0, "width": 800, "height": 600}},
        "b": {"id": "b", "type": "shape", "shape": "rectangle",
              "frame": {"x": 100, "y": 150, "width": 200, "height": 300}},
    }
    doc = DesignDocument.model_validate({
        "schemaVersion": 2, "id": "d_two", "title": "Two", "revision": 1,
        "updatedAt": "2026-10-18T00:00:00+00:00", "tokens": wire(default_tokens()),
        "pages": [
            {"id": "p1", "artboards": [{"id": "wide", "width": 800, "height": 600, "nodeIds": ["a"]}]},
            {"id": "p2", "artboards": [{"id": "tall", "width": 400, "height": 1200, "nodeIds": ["b"]}]},
        ],
        "nodes": nodes, "comments": [],
    })
    prs = _slides(export_document(doc, "pptx").data)
    assert len(prs.slides) == 2
    second = _by_name(prs.slides[1])
    # 400x1200 on an 800x600 slide: scale 0.5 of 96 dpi, centred horizontally.
    scale = 600 * EMU / 1200
    left = (800 * EMU - 400 * scale) / 2
    board = second["artboard tall"]
    assert (board.left, board.top) == (round(left), 0)
    assert (board.width, board.height) == (round(400 * scale), 600 * EMU)
    b = second["b"]
    assert (b.left, b.top) == (round(left + 100 * scale), round(150 * scale))
    assert (b.width, b.height) == (round(200 * scale), round(300 * scale))


def test_pptx_keeps_a_huge_artboard_inside_powerpoints_limit() -> None:
    pytest.importorskip("pptx", reason="python-pptx is the documents-out extra")
    nodes = {"r": {"id": "r", "type": "shape", "shape": "rectangle",
                   "frame": {"x": 5000, "y": 0, "width": 1000, "height": 100}}}
    prs = _slides(export_document(
        _doc([{"id": "a1", "width": 10_000, "height": 1_000, "nodeIds": ["r"]}], nodes), "pptx"
    ).data)
    assert prs.slide_width == export_pptx.MAX_SLIDE_EMU
    scale = export_pptx.MAX_SLIDE_EMU / 10_000
    assert prs.slide_height == round(1_000 * scale)
    r = _by_name(prs.slides[0])["r"]
    assert (r.left, r.width) == (round(5000 * scale), round(1000 * scale))


def test_pptx_embeds_a_png_whose_bytes_are_passed_and_only_then() -> None:
    pptx = pytest.importorskip("pptx", reason="python-pptx is the documents-out extra")
    from tests.test_design_export_is_inert_and_checks_tokens import _png

    png = _png(4, 4)
    nodes = {"i": {"id": "i", "type": "image", "frame": {"x": 10, "y": 10, "width": 50, "height": 50},
                   "assetId": "logo", "alt": "Logo"}}
    doc = _doc([{"id": "a1", "width": 100, "height": 100, "nodeIds": ["i"]}], nodes)
    with_bytes = _by_name(_slides(export_pptx.render_pptx(doc, assets={"logo": png})).slides[0])
    assert with_bytes["i"].shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE
    without = _by_name(_slides(export_pptx.render_pptx(doc)).slides[0])
    assert without["i"].shape_type != pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE


def test_html_text_is_sized_by_the_slide_scale_and_keeps_alignment_and_family(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    pytest.importorskip("pptx", reason="python-pptx is the documents-out extra")
    from pptx.enum.text import PP_ALIGN

    # The first artboard sets the slide (300x160 px); the 600x320 one is scaled to half of it.
    scaled = _FakeTextBox("Scaled", {"x": 10, "y": 10, "width": 200, "height": 40},
                          "font-size:32px;font-weight:400;color:rgb(0, 0, 0);text-align:right;"
                          "font-family:'Georgia', serif")
    _fake_html_artboard(monkeypatch, [scaled], size=(600, 320))
    doc = _doc([
        {"id": "v1", "width": 300, "height": 160, "nodeIds": []},
        {"id": "html", "width": 600, "height": 320, "nodeIds": [], "html": "<main><p>Scaled</p></main>"},
    ], {})
    presentation = _slides(export_document(doc, "pptx").data)
    (shape,) = [s for s in presentation.slides[1].shapes if s.has_text_frame and s.text == "Scaled"]
    paragraph = shape.text_frame.paragraphs[0]
    assert paragraph.runs[0].font.size.pt == pytest.approx(12)  # 32 px times the 0.5 slide scale
    assert paragraph.alignment == PP_ALIGN.RIGHT
    assert paragraph.runs[0].font.name == "Georgia"


def test_a_transparent_text_keeps_the_default_colour_and_one_bad_element_does_not_stop_the_export(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    pytest.importorskip("pptx", reason="python-pptx is the documents-out extra")
    _fake_html_artboard(monkeypatch, [
        _FakeTextBox("Hidden", {"x": 12, "y": 18, "width": 80, "height": 20},
                     "font-size:20px;font-weight:400;color:rgba(255, 0, 0, 0)"),
        _FakeTextBox("", {"x": 0, "y": 0, "width": 1, "height": 1}, "",
                     error=RuntimeError("Node is not an HTMLElement")),
        _FakeTextBox("Still editable", {"x": 12, "y": 60, "width": 120, "height": 20},
                     "font-size:16px;font-weight:400;color:rgb(0, 0, 0)"),
    ])
    doc = _doc([{"id": "html", "width": 300, "height": 160, "nodeIds": [],
                 "html": "<main><p>Still editable</p></main>"}], {})
    presentation = _slides(export_document(doc, "pptx").data)
    shapes = [s for s in presentation.slides[0].shapes if s.has_text_frame]
    texts = {s.text: s for s in shapes}
    assert "Still editable" in texts
    hidden = texts["Hidden"].text_frame.paragraphs[0].runs[0].font
    assert hidden.color.type is None  # alpha 0: nothing to draw, so no colour is set
