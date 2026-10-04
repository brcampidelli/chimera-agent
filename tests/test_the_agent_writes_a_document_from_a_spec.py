"""create_document (study 29, P6.2): Word, Excel, PowerPoint and PDF from a declarative spec.

Every test here OPENS the file the tool wrote and reads back the paragraph, the cell, the slide. A
writer test that checks only "a file exists and the tool said saved" passes for an empty document,
and this project has paid for that family of test before ("nothing gives an error").
"""

from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

import pytest

from chimera.tools import create_document as cd
from chimera.tools.document_spec import SpecError, markdown_to_spec, parse_spec
from chimera.tools.workspace import PathEscapesWorkspaceError
from chimera.tools.write_region import WriteRegion

docx = pytest.importorskip("docx")  # the dev extra installs documents-out; see pyproject.toml
openpyxl = pytest.importorskip("openpyxl")
pptx = pytest.importorskip("pptx")

REPORT = {
    "title": "Relatório do trimestre",
    "blocks": [
        {"type": "heading", "text": "Resumo", "level": 1},
        {"type": "paragraph", "text": "As vendas subiram (12%) no período."},
        {"type": "bullets", "items": ["Primeiro ponto", "Segundo ponto"], "numbered": True},
        {"type": "table", "name": "Vendas", "columns": ["Região", "Total"], "rows": [["Sul", 10], ["Norte", 2.5]]},
        {"type": "page_break"},
        {"type": "heading", "text": "Anexo", "level": 2},
    ],
}


def _tool(tmp_path: Path, region: WriteRegion | None = None) -> cd.CreateDocumentTool:
    return cd.CreateDocumentTool(tmp_path, write_region=region)


# --- each format, read back ------------------------------------------------------------------------


def test_a_word_document_holds_the_headings_paragraphs_list_and_table_it_was_given(tmp_path: Path) -> None:
    answer = _tool(tmp_path).run(spec=REPORT, path="report.docx")
    assert answer.startswith("saved docx document"), answer

    doc = docx.Document(str(tmp_path / "report.docx"))
    texts = [p.text for p in doc.paragraphs]
    styles = {p.text: p.style.name for p in doc.paragraphs}
    assert texts[0] == "Relatório do trimestre" and styles["Relatório do trimestre"] == "Title"
    assert styles["Resumo"] == "Heading 1" and styles["Anexo"] == "Heading 2"
    assert "As vendas subiram (12%) no período." in texts
    assert styles["Primeiro ponto"] == "List Number" and "Segundo ponto" in texts
    table = doc.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["Região", "Total"]
    assert [c.text for c in table.rows[2].cells] == ["Norte", "2.5"]
    assert doc.core_properties.title == "Relatório do trimestre"


def test_a_workbook_keeps_numbers_as_numbers_and_names_the_sheet_after_the_table(tmp_path: Path) -> None:
    answer = _tool(tmp_path).run(spec=REPORT, path="report.xlsx")
    assert "1 sheet(s), 2 row(s)" in answer

    wb = openpyxl.load_workbook(tmp_path / "report.xlsx")
    assert wb.sheetnames == ["Vendas"]
    ws = wb["Vendas"]
    assert [c.value for c in ws[1]] == ["Região", "Total"] and ws["A1"].font.b
    assert ws["A2"].value == "Sul" and ws["B2"].value == 10 and ws["B3"].value == 2.5


@pytest.mark.parametrize("payload", ["=1+1", "=HYPERLINK(\"http://x\",\"clique\")", "+cmd|' /C calc'!A0", "-2+3", "@SUM(A1)"])
def test_model_text_in_a_cell_never_becomes_a_formula(tmp_path: Path, payload: str) -> None:
    """Formula injection: a cell that opens as `=WEBSERVICE(...)` runs on the owner's machine. The
    text may have come from a fetched page, so it is stored as text and marked with Excel's quote
    prefix — read back from the XML itself, not only through openpyxl."""
    spec = {"sheets": [{"name": "Dados", "columns": ["v"], "rows": [[payload]]}]}
    assert _tool(tmp_path).run(spec=spec, path="x.xlsx").startswith("saved xlsx")

    with zipfile.ZipFile(tmp_path / "x.xlsx") as z:
        sheet_xml = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert "<f>" not in sheet_xml and "<f " not in sheet_xml
    cell = openpyxl.load_workbook(tmp_path / "x.xlsx")["Dados"]["A2"]
    assert cell.value == payload and cell.data_type == "s" and cell.quotePrefix


def test_a_sheet_name_excel_would_refuse_is_made_legal_and_kept_unique(tmp_path: Path) -> None:
    spec = {"sheets": [
        {"name": "a/b:c*?[x]", "rows": [[1]]},
        {"name": "A_B_C___X_", "rows": [[2]]},  # the same name to Excel, which ignores case
        {"name": "x" * 40, "rows": [[3]]},
        {"name": "", "rows": [[4]]},
    ]}
    assert _tool(tmp_path).run(spec=spec, path="s.xlsx").startswith("saved")
    names = openpyxl.load_workbook(tmp_path / "s.xlsx").sheetnames
    assert names[:2] == ["a_b_c___x_", "A_B_C___X_ (2)"]
    assert all(len(n) <= 31 for n in names) and names[3] == "Sheet4" and len(set(names)) == 4


def test_a_deck_has_its_cover_bullets_notes_and_table(tmp_path: Path) -> None:
    spec = {
        "title": "Plano 2027",
        "slides": [
            {"title": "Metas", "bullets": ["Crescer", "Contratar"], "notes": "falar devagar"},
            {"title": "Números", "table": {"columns": ["Ano", "Meta"], "rows": [[2027, "10k"]]}},
        ],
    }
    assert "2 slide(s)" in _tool(tmp_path).run(spec=spec, path="plano.pptx")

    deck = pptx.Presentation(str(tmp_path / "plano.pptx"))
    slides = list(deck.slides)
    assert [s.shapes.title.text for s in slides] == ["Plano 2027", "Metas", "Números"]
    body = slides[1].placeholders[1].text_frame
    assert [p.text for p in body.paragraphs] == ["Crescer", "Contratar"]
    assert slides[1].notes_slide.notes_text_frame.text == "falar devagar"
    table = next(sh for sh in slides[2].shapes if sh.has_table).table
    assert table.cell(1, 0).text == "2027" and table.cell(1, 1).text == "10k"


def test_a_long_word_table_is_built_in_linear_time(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    # `grid.cell(r, c)` rebuilds the list of every cell in the table on each call: 400x10 took 154 s
    # and the spec allows 5000x50. Counted rather than only timed, so a slow machine cannot hide it.
    import time

    from docx.table import Table as DocxTable

    real = DocxTable.__dict__["_cells"]  # the property object itself, not its value
    calls = [0]

    def counted(self: DocxTable) -> object:
        calls[0] += 1
        if calls[0] > 50:  # stop a quadratic build here instead of after twenty minutes
            raise AssertionError("the table's full cell list was rebuilt per cell")
        return real.fget(self)

    monkeypatch.setattr(DocxTable, "_cells", property(counted))
    rows = [[f"r{r}c{c}" for c in range(10)] for r in range(1000)]
    spec = {"blocks": [{"type": "table", "columns": [f"c{i}" for i in range(10)], "rows": rows}]}
    started = time.perf_counter()
    assert _tool(tmp_path).run(spec=spec, path="long.docx").startswith("saved docx")
    assert time.perf_counter() - started < 20  # measured 1.4 s; the old loop needed ~20 minutes
    assert calls[0] <= 2, calls[0]

    table = docx.Document(str(tmp_path / "long.docx")).tables[0]
    assert len(table.rows) == 1001 and table.rows[1000].cells[9].text == "r999c9"
    assert table.rows[0].cells[3].paragraphs[0].runs[0].bold


# --- PDF: no extra, read back with nothing but the standard library --------------------------------


def _pdf_text(data: bytes) -> list[str]:
    """Every string a `Tj` draws, unescaped and decoded from WinAnsi."""
    out: list[str] = []
    for raw in re.findall(rb"\(((?:[^()\\]|\\.)*)\) Tj", data):
        decoded = re.sub(rb"\\([0-7]{3})", lambda m: bytes([int(m.group(1), 8)]), raw)
        decoded = re.sub(rb"\\(.)", rb"\1", decoded)
        out.append(decoded.decode("cp1252"))
    return out


def _check_xref(data: bytes) -> int:
    """Every offset in the cross-reference table points at the object it names; returns the count."""
    start = int(data.rsplit(b"startxref\n", 1)[1].split(b"\n", 1)[0])
    assert data[start:].startswith(b"xref\n")
    lines = data[start:].split(b"\n")
    count = int(lines[1].split()[1])
    for number in range(1, count):
        offset = int(lines[2 + number].split()[0])
        assert data[offset:].startswith(b"%d 0 obj" % number), number
    return count


def test_a_pdf_is_written_with_no_extra_and_its_text_and_structure_read_back(tmp_path: Path) -> None:
    answer = _tool(tmp_path).run(spec=REPORT, path="report.pdf")
    assert answer.startswith("saved pdf document") and "?" not in answer.split(")")[-1]

    data = (tmp_path / "report.pdf").read_bytes()
    assert data.startswith(b"%PDF-1.4") and data.rstrip().endswith(b"%%EOF")
    assert _check_xref(data) > 6
    text = _pdf_text(data)
    assert "Relatório do trimestre" in text and "Resumo" in text
    assert "As vendas subiram (12%) no período." in text  # the parentheses survived escaping
    assert {"Região", "Total", "Sul", "10", "Norte", "2.5"} <= set(text)
    assert data.count(b"/Type /Page ") == 2  # the page break made a second page


def test_a_pdf_says_how_many_characters_its_font_could_not_draw(tmp_path: Path) -> None:
    spec = {"blocks": [{"type": "paragraph", "text": "olá 世界"}]}
    answer = _tool(tmp_path).run(spec=spec, path="x.pdf")
    assert "2 character(s) have no glyph" in answer
    assert "olá ??" in _pdf_text((tmp_path / "x.pdf").read_bytes())


def test_a_long_table_continues_on_new_pages_with_its_header_repeated(tmp_path: Path) -> None:
    rows = [[f"linha {i}", i] for i in range(120)]
    spec = {"blocks": [{"type": "table", "columns": ["Nome", "N"], "rows": rows}]}
    _tool(tmp_path).run(spec=spec, path="long.pdf")
    data = (tmp_path / "long.pdf").read_bytes()
    pages = data.count(b"/Type /Page ")
    text = _pdf_text(data)
    assert pages >= 2 and text.count("Nome") == pages and "linha 119" in text
    _check_xref(data)


def test_a_table_cell_taller_than_a_page_continues_on_the_next_one_with_every_word(tmp_path: Path) -> None:
    # The adversarial review's case: 1500 words in one cell. The row used to be cut to one page and
    # 300 of 1500 words reached the PDF while the answer said "saved". Now the row continues.
    words = " ".join(f"palavra{i}" for i in range(1500))
    spec = {"blocks": [{"type": "table", "columns": ["Nota", "N"], "rows": [[words, 1], ["depois", 2]]}]}
    answer = _tool(tmp_path).run(spec=spec, path="tall.pdf")
    assert answer.startswith("saved pdf document")
    data = (tmp_path / "tall.pdf").read_bytes()
    drawn = " ".join(_pdf_text(data)).split()
    assert [w for w in drawn if w.startswith("palavra")] == words.split()  # all of them, in order
    pages = data.count(b"/Type /Page ")
    assert pages >= 3 and drawn.count("Nota") == pages and "depois" in drawn
    _check_xref(data)


def test_a_narrow_column_of_long_text_loses_no_character(tmp_path: Path) -> None:
    # Fifty columns leave ~10pt a column, so nearly every character is a line of its own: the case
    # where the old one-page cut lost the most with the least text.
    text = "abcdefghij" * 40
    row = [text] + ["x"] * 49
    spec = {"blocks": [{"type": "table", "rows": [row]}]}
    _tool(tmp_path).run(spec=spec, path="narrow.pdf")
    drawn = _pdf_text((tmp_path / "narrow.pdf").read_bytes())
    assert "".join(s for s in drawn if s != "x") == text


def test_no_table_text_is_drawn_below_the_bottom_margin_of_a_page(tmp_path: Path) -> None:
    # The review's second case: a tall header in narrow columns, then a tall row. The row began below
    # the repeated header with no room left and was drawn at y < 0, off the page — invisible, while
    # the tool said "saved". Every text position must stay inside the margins.
    from chimera.tools.pdf_writer import MARGIN

    columns = [f"coluna_numero_{i:03d}" for i in range(50)]  # 17 characters, one per line at ~10pt
    rows = [["kq " * 120] + ["y"] * 49 for _ in range(3)]  # k and q appear nowhere else
    spec = {"blocks": [{"type": "paragraph", "text": "antes " * 400}, {"type": "table", "columns": columns, "rows": rows}]}
    _tool(tmp_path).run(spec=spec, path="margin.pdf")
    data = (tmp_path / "margin.pdf").read_bytes()
    ys = [float(m) for m in re.findall(rb"[\d.]+ (-?[\d.]+) Td", data)]
    assert ys and min(ys) >= MARGIN - 1e-6, min(ys)
    drawn = "".join(_pdf_text(data))
    assert drawn.count("k") == drawn.count("q") == 3 * 120  # and none of it was cut
    _check_xref(data)


# --- refusals --------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("spec", "path", "fmt", "says"),
    [
        ({}, "x.docx", None, "nothing to write"),
        ("{not json", "x.docx", None, "not valid JSON"),
        ({"blocks": [{"type": "script", "text": "x"}]}, "x.docx", None, "'script' is not one of"),
        ({"blocks": [{"type": "heading", "text": "x", "level": 7}]}, "x.docx", None, "level must be"),
        ({"blocks": [{"type": "paragraph", "text": {"a": 1}}]}, "x.docx", None, "must be text"),
        ({"sheets": [{"rows": [[{"nested": 1}]]}]}, "x.xlsx", None, "rows[0][0] must be text"),
        ({"blocks": [{"type": "paragraph", "text": "x"}]}, "x.docx", "pdf", "does not match"),
        ({"blocks": [{"type": "paragraph", "text": "x"}]}, "x.txt", None, "say which format"),
        ({"blocks": [{"type": "paragraph", "text": "x"}]}, "x.docx", "exe", "unknown format"),
        ({"blocks": [{"type": "paragraph", "text": "x"}]}, "x.xlsx", None, "needs 'sheets'"),
        ({"blocks": [{"type": "paragraph", "text": "x"}]}, "x.pptx", None, "needs 'slides'"),
    ],
)
def test_a_spec_it_cannot_render_is_refused_and_nothing_is_written(
    tmp_path: Path, spec: object, path: str, fmt: str | None, says: str
) -> None:
    answer = _tool(tmp_path).run(spec=spec, path=path, format=fmt)
    assert answer.startswith("error: invalid document spec") and says in answer, answer
    assert list(tmp_path.iterdir()) == []


@pytest.mark.parametrize("path", ["README.md", "src/app.py", "notes.txt"])
def test_a_format_never_overwrites_a_file_whose_extension_names_something_else(tmp_path: Path, path: str) -> None:
    # The review's case: path='README.md', format='pdf' answered "saved pdf document" and the README
    # began with %PDF-1.4 — a text file replaced by binary, with no diff for the ledger to show.
    target = tmp_path / path
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text("# o README do dono\n", encoding="utf-8")
    answer = _tool(tmp_path).run(spec={"blocks": [{"type": "paragraph", "text": "x"}]}, path=path, format="pdf")
    assert answer.startswith("error: invalid document spec") and "ending in .pdf" in answer, answer
    assert target.read_text(encoding="utf-8") == "# o README do dono\n"


def test_a_path_with_no_extension_gets_the_formats(tmp_path: Path) -> None:
    (tmp_path / "LICENSE").write_text("MIT\n", encoding="utf-8")
    answer = _tool(tmp_path).run(spec={"blocks": [{"type": "paragraph", "text": "x"}]}, path="LICENSE", format="pdf")
    assert answer.startswith("saved pdf document") and "LICENSE.pdf" in answer, answer
    assert (tmp_path / "LICENSE").read_text(encoding="utf-8") == "MIT\n"
    assert (tmp_path / "LICENSE.pdf").read_bytes().startswith(b"%PDF-")


def test_a_runaway_table_is_refused_before_it_renders(tmp_path: Path) -> None:
    spec = {"sheets": [{"rows": [[1]] * 5001}]}
    assert "limit is 5000" in _tool(tmp_path).run(spec=spec, path="big.xlsx")


def test_a_path_outside_the_workspace_is_refused(tmp_path: Path) -> None:
    ws = tmp_path / "ws"
    ws.mkdir()
    with pytest.raises(PathEscapesWorkspaceError):
        _tool(ws).run(spec=REPORT, path=str(tmp_path / "escaped.pdf"))
    assert not (tmp_path / "escaped.pdf").exists()


@pytest.mark.parametrize("target", [".git/report.pdf", ".chimera/report.pdf"])
def test_the_machinery_the_workspace_is_governed_by_is_never_written(tmp_path: Path, target: str) -> None:
    assert _tool(tmp_path).run(spec=REPORT, path=target).startswith("error:")
    assert not (tmp_path / target).exists()


def test_a_declared_write_region_holds_for_documents_too(tmp_path: Path) -> None:
    region = WriteRegion(["out/**"], tmp_path)
    assert _tool(tmp_path, region).run(spec=REPORT, path="report.pdf").startswith("error:")
    assert _tool(tmp_path, region).run(spec=REPORT, path="out/report.pdf").startswith("saved")


def test_without_the_extra_the_answer_names_it_and_leaves_no_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def missing(_spec: object) -> bytes:
        raise ImportError("No module named 'docx'")

    monkeypatch.setitem(cd._RENDERERS, "docx", missing)
    answer = _tool(tmp_path).run(spec=REPORT, path="report.docx")
    assert "documents-out" in answer and answer.startswith("error:")
    assert not (tmp_path / "report.docx").exists()


def test_a_render_that_fails_half_way_leaves_no_truncated_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def broken(_spec: object) -> bytes:
        raise RuntimeError("disk of the renderer exploded")

    monkeypatch.setitem(cd._RENDERERS, "docx", broken)
    assert "could not render" in _tool(tmp_path).run(spec=REPORT, path="report.docx")
    assert not (tmp_path / "report.docx").exists()


def test_control_characters_xml_cannot_hold_are_dropped_instead_of_failing_the_document(tmp_path: Path) -> None:
    spec = {"blocks": [{"type": "paragraph", "text": "a\x00b\x07c\td"}]}
    assert _tool(tmp_path).run(spec=spec, path="c.docx").startswith("saved")
    assert docx.Document(str(tmp_path / "c.docx")).paragraphs[0].text == "abc\td"


# --- where it is offered ---------------------------------------------------------------------------


def _registry_names(monkeypatch: pytest.MonkeyPatch, tmp_path: Path, value: str | None) -> set[str]:
    from chimera.config import get_settings
    from chimera.tools.builtin import default_registry

    if value is None:
        monkeypatch.delenv("CHIMERA_CREATE_DOCUMENT", raising=False)
    else:
        monkeypatch.setenv("CHIMERA_CREATE_DOCUMENT", value)
    monkeypatch.setenv("CHIMERA_HOME", str(tmp_path / "home"))
    get_settings.cache_clear()
    try:
        return set(default_registry(tmp_path, host_exec_confirm=None).names())
    finally:
        get_settings.cache_clear()


def test_the_tool_is_off_until_the_owner_switches_it_on(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    assert "create_document" not in _registry_names(monkeypatch, tmp_path, None)
    assert "create_document" in _registry_names(monkeypatch, tmp_path, "1")


def test_the_tools_screen_offers_the_switch_for_it() -> None:
    from chimera.api.config_api import ALLOWED_KEYS
    from chimera.tools.conditional import CONDITIONAL_TOOLS

    row = next(t for t in CONDITIONAL_TOOLS if t.name == "create_document")
    assert row.switchable and not row.default_on and row.variables == ("CHIMERA_CREATE_DOCUMENT",)
    assert "CHIMERA_CREATE_DOCUMENT" in ALLOWED_KEYS


def test_a_read_only_posture_denies_it_like_any_other_writer() -> None:
    from chimera.api.posture import Posture, resolve

    assert "create_document" in resolve(Posture(reach="read_only", approval="never")).deny_tools


# --- the deliver path: Markdown from the model, the same renderers -------------------------------

_MARKDOWN = """# Relatório

Texto **inicial** do relatório.

## Números

### Vendas
| Região | Total |
|---|---:|
| Sul | =1+1 |
| Norte | 7 |

- um
- dois

1. primeiro
2. segundo
"""


def test_markdown_becomes_headings_lists_and_named_tables() -> None:
    spec = parse_spec(markdown_to_spec(_MARKDOWN))
    kinds = [type(b).__name__ for b in spec.blocks]
    assert kinds == ["Heading", "Paragraph", "Heading", "Heading", "Table", "Bullets", "Bullets"]
    table = spec.blocks[4]
    assert table.name == "Vendas" and table.columns == ("Região", "Total")  # type: ignore[union-attr]
    assert spec.blocks[1].text == "Texto inicial do relatório."  # type: ignore[union-attr]
    assert spec.blocks[6].numbered and not spec.blocks[5].numbered  # type: ignore[union-attr]


def test_deliver_builds_a_word_file_from_the_models_markdown() -> None:
    from chimera.deliver import render_deliverable

    data, _ = render_deliverable(_MARKDOWN, "docx")
    doc = docx.Document(io.BytesIO(data))
    assert doc.paragraphs[0].text == "Relatório" and doc.paragraphs[0].style.name == "Title"
    assert doc.tables[0].rows[1].cells[1].text == "=1+1"


def test_deliver_builds_a_workbook_whose_model_text_stays_text() -> None:
    from chimera.deliver import render_deliverable

    data, _ = render_deliverable(_MARKDOWN, "xlsx")
    ws = openpyxl.load_workbook(io.BytesIO(data))["Vendas"]
    assert ws["B2"].value == "=1+1" and ws["B2"].data_type == "s"


def test_a_delivered_workbook_holds_plain_numbers_as_numbers_so_they_sum() -> None:
    # Every Markdown cell used to arrive as text with the quote prefix: `| jan | 1200.50 |` read back
    # as ('1200.50', 's'), and SUM and charts saw nothing. A plain number now is one; anything that
    # could lead a formula, a padded id, a long account number or a separated thousand stays text.
    from chimera.deliver import render_deliverable

    markdown = (
        "## Caixa\n| mês | valor | id | conta | obs |\n|---|---|---|---|---|\n"
        "| jan | 1200.50 | 007 | 1234567890123456 | =SUM(B2:B3) |\n"
        "| fev | -35 | 0 | 12 | 1,200.50 |\n"
    )
    ws = openpyxl.load_workbook(io.BytesIO(render_deliverable(markdown, "xlsx")[0]))["Caixa"]
    assert ws["B2"].value == 1200.5 and ws["B2"].data_type == "n"
    assert ws["B3"].value == -35 and ws["B3"].data_type == "n" and ws["C3"].value == 0
    for ref, text in (("C2", "007"), ("D2", "1234567890123456"), ("E2", "=SUM(B2:B3)"), ("E3", "1,200.50")):
        assert ws[ref].value == text and ws[ref].data_type == "s", ref
    assert ws["A1"].value == "mês" and ws["A2"].value == "jan"


def test_a_delivered_word_table_prints_a_number_as_the_model_wrote_it() -> None:
    from chimera.deliver import render_deliverable

    data, _ = render_deliverable("| mês | valor |\n|---|---|\n| jan | 1200.50 |\n", "docx")
    assert docx.Document(io.BytesIO(data)).tables[0].rows[1].cells[1].text == "1200.50"


def test_deliver_refuses_an_xlsx_when_the_answer_has_no_table() -> None:
    from chimera.deliver import render_deliverable

    with pytest.raises(ValueError, match="sheets"):
        render_deliverable("# Só texto\n\nnenhuma tabela", "xlsx")
    with pytest.raises(SpecError):
        render_deliverable("", "pdf")


def test_deliver_asks_for_a_path_before_it_spends_anything_on_a_binary_format(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from typer.testing import CliRunner

    from chimera.cli.main import app

    called: list[str] = []
    monkeypatch.setattr("chimera.deliver.produce_deliverable", lambda *a, **k: called.append("x") or "")
    result = CliRunner().invoke(app, ["deliver", "a report", "--format", "docx"])
    assert result.exit_code == 2 and "--out" in result.output and called == []
    unknown = CliRunner().invoke(app, ["deliver", "a report", "--format", "rtf"])
    assert unknown.exit_code == 2 and called == []


def test_a_deliverable_that_cannot_convert_keeps_its_markdown_without_overwriting_the_owners(
    tmp_path: Path,
) -> None:
    # `deliver "plano" -f xlsx -o plano.xlsx` with no table in the answer saved the Markdown to
    # plano.md — over the owner's own plano.md, which the command never named.
    import typer

    from chimera.cli.main import _write_binary_deliverable

    (tmp_path / "plano.md").write_text("o plano do dono\n", encoding="utf-8")
    (tmp_path / "plano-1.md").write_text("outra versão do dono\n", encoding="utf-8")
    with pytest.raises(typer.Exit):
        _write_binary_deliverable("# Plano\n\nsó texto, nenhuma tabela\n", "xlsx", tmp_path / "plano.xlsx")
    assert (tmp_path / "plano.md").read_text(encoding="utf-8") == "o plano do dono\n"
    assert (tmp_path / "plano-1.md").read_text(encoding="utf-8") == "outra versão do dono\n"
    assert (tmp_path / "plano-2.md").read_text(encoding="utf-8").startswith("# Plano")
    assert not (tmp_path / "plano.xlsx").exists()


def test_with_no_markdown_beside_it_the_fallback_is_the_plain_name(tmp_path: Path) -> None:
    import typer

    from chimera.cli.main import _write_binary_deliverable

    with pytest.raises(typer.Exit):
        _write_binary_deliverable("# Plano\n\nsó texto\n", "xlsx", tmp_path / "plano.xlsx")
    assert (tmp_path / "plano.md").read_text(encoding="utf-8").startswith("# Plano")
